#!/usr/bin/env python3
"""KSC 15.1 hardware CPU inventory joined to an existing CAD CSV.

Read-only KlAkOAPI calls. Pilot: SSH loopback tunnel; production: valid TLS CA.
Derived from a real inventory of 104 computers / 103 CPU host IDs.
This sanitized published revision needs its own end-to-end validation.
"""
import argparse
import csv
import getpass
import re
import warnings
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from KlAkOAPI.AdmServer import KlAkAdmServer
from KlAkOAPI.ChunkAccessor import KlAkChunkAccessor
from KlAkOAPI.HostGroup import KlAkHostGroup
from KlAkOAPI.Params import KlAkArray
from KlAkOAPI.SrvView import KlAkSrvView
from ksc_hardware_rules import (
    classify_storage, host_status, modernization_priority, vram_mb,
    normalized_cpu_cores,
)


def get(item, name):
    if hasattr(item, "GetNames"):
        return item[name] if name in item.GetNames() else None
    if isinstance(item, dict):
        return item.get(name)
    return None


def string(item, name):
    value = get(item, name)
    return "" if value is None else str(value).strip()


def integer(item, field):
    try:
        v = get(item, field)
        n = int(v)
        return n if n >= 0 else None
    except (TypeError, ValueError, OverflowError):
        return None


def vendor(model):
    text = model.casefold()
    if "intel" in text:
        return "Intel"
    if any(part in text for part in ("amd", "ryzen", "epyc", "athlon")):
        return "AMD"
    return "Не определён"


def fetch_hosts(server):
    api = KlAkHostGroup(server)
    chunks = KlAkChunkAccessor(server)
    answer = api.FindHosts(
        '(KLHST_WKS_HOSTNAME="*")',
        ["KLHST_WKS_HOSTNAME", "KLHST_WKS_DN", "HST_VM_TYPE", "KLHST_WKS_LAST_NAGENT_CONNECTED"],
        [], {"KLGRP_FIND_FROM_CUR_VS_ONLY": True}, 3600,
    )
    accessor = answer.OutPar("strAccessor")
    found = {}
    try:
        count = int(chunks.GetItemsCount(accessor).RetVal())
        for start in range(0, count, 100):
            response = chunks.GetItemsChunk(accessor, start, min(100, count - start))
            records = response.OutPar("pChunk")["KLCSP_ITERATOR_ARRAY"]
            for record in records:
                host_id = string(record, "KLHST_WKS_HOSTNAME")
                computer = string(record, "KLHST_WKS_DN")
                if host_id:
                    found[host_id] = {"name": (computer or host_id).upper(), "vm_type": get(record, "HST_VM_TYPE"), "last_seen": string(record, "KLHST_WKS_LAST_NAGENT_CONNECTED")}
    finally:
        chunks.Release(accessor)
    return found


def fetch_hardware(server):
    api = KlAkSrvView(server)
    requested = KlAkArray([
        "KLHST_WKS_HOSTNAME", "dev_type", "dev_id", "dev_name", "dev_descr",
        "dev_manuf", "dev_xpar2", "dev_xpar3", "CpuCores", "CpuThreads",
        "CpuSpeed", "StorageRotationRate", "VideoCardMemorySize",
        "RamType", "RamPartNum", "SerialNum",
    ])
    response = api.ResetIterator("HWInvPCSrvViewName", "", requested, KlAkArray([]), {}, lifetimeSec=3600)
    token = response.OutPar("wstrIteratorId")
    by_host, types = defaultdict(list), Counter()
    try:
        total = int(api.GetRecordCount(token).RetVal())
        for offset in range(0, total, 200):
            data = api.GetRecordRange(token, offset, min(total, offset + 200))
            for item in data.OutPar("pRecords")["KLCSP_ITERATOR_ARRAY"]:
                kind = integer(item, "dev_type")
                types[kind] += 1
                host_id = string(item, "KLHST_WKS_HOSTNAME")
                if not host_id:
                    continue
                by_host[host_id].append({
                    "type": kind,
                    "device_id": string(item, "dev_id"),
                    "serial": string(item, "SerialNum"),
                    "model": string(item, "dev_name") or string(item, "dev_descr"),
                    "manufacturer": string(item, "dev_manuf"),
                    "cores": integer(item, "CpuCores"),
                    "threads": integer(item, "CpuThreads"),
                    "speed": integer(item, "CpuSpeed"),
                    "xpar2": integer(item, "dev_xpar2"),
                    "xpar3": integer(item, "dev_xpar3"),
                    "rpm": integer(item, "StorageRotationRate"),
                    "vram": integer(item, "VideoCardMemorySize"),
                    "ram_type": string(item, "RamType"),
                    "ram_part": string(item, "RamPartNum"),
                })
    finally:
        api.ReleaseIterator(token)
    return by_host, total, types


def gib(value):
    return round(value / (1024 ** 3), 2) if value else None


def type_of_storage(item):
    return classify_storage(item["model"], item["rpm"])[0]


def dedup_hardware(records, category):
    used, unique, duplicates = set(), [], 0
    for item in records:
        if item["type"] != category:
            continue
        if category == 1:
            # Identical CPU signatures are a duplicate unless independently verified.
            key = (item["model"].casefold(), item["cores"], item["threads"], item["speed"])
        else:
            # Do NOT merge equal-size/equal-model DIMMs or SSDs: they can be real devices.
            key = (item["device_id"] or item["serial"])
            if not key:
                unique.append(item)
                continue
        if key in used:
            duplicates += 1
        else:
            used.add(key)
            unique.append(item)
    return unique, duplicates



def hardware_info(records):
    cpus, duplicates = dedup_hardware(records, 1)
    rams, _ = dedup_hardware(records, 2)
    disks, _ = dedup_hardware(records, 3)
    gpus, _ = dedup_hardware(records, 4)
    names = "; ".join(sorted({r["model"] for r in cpus}))
    _cores_raw, cores, core_audit = normalized_cpu_cores(cpus)
    threads = sum(r["threads"] for r in cpus) if cpus and all(r["threads"] for r in cpus) else ""
    freq = "; ".join(sorted({str(r["speed"]) for r in cpus if r["speed"]}))
    ram = gib(sum(r["xpar3"] for r in rams)) if rams and all(r["xpar3"] for r in rams) else ""
    gpu = "; ".join(sorted({r["model"] for r in gpus}))
    vram = "; ".join(sorted({str(vram_mb(r["vram"])) for r in gpus if vram_mb(r["vram"]) is not None}))
    disk_gb = round(sum(gib(r["xpar2"]) for r in disks), 2) if disks and all(r["xpar2"] for r in disks) else ""
    disk_type = "; ".join(sorted({type_of_storage(r) for r in disks}))
    disk_models = "; ".join(sorted({r["model"] for r in disks}))
    checks = (["Повторы CPU удалены"] if duplicates else [])
    if core_audit:
        checks.append(core_audit)
    checks += ["Нет CPU"] if not cpus else []
    checks += ["Нет RAM"] if not rams else []
    checks += ["Нет GPU"] if not gpus else []
    checks += ["Нет дисков"] if not disks else []
    data = [names or "Нет данных", vendor(names), len([r for r in records if r["type"] == 1]),
            len(cpus), cores, threads, freq, ram, len(rams), gpu, vram,
            disk_gb, len(disks), disk_type, disk_models, "; ".join(checks) or "OK"]
    return data, cpus, rams, gpus, disks, duplicates

def read_cad(path):
    if not path.is_file():
        raise RuntimeError(f"CAD CSV is missing: {path}")
    by_pc = defaultdict(lambda: defaultdict(set))
    records = 0
    with path.open(encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source, delimiter=";")
        expected = {"Computer", "Software", "Release", "ServicePack", "DisplayVersion"}
        if not expected.issubset(reader.fieldnames or []):
            raise RuntimeError(f"Unexpected CAD CSV columns: {reader.fieldnames}")
        for item in reader:
            pc = (item["Computer"] or "").strip().upper()
            app = (item["Software"] or "").strip()
            release = " ".join(v.strip() for v in (item["Release"], item["ServicePack"]) if v and v.strip())
            if not release:
                release = ("build " + (item["DisplayVersion"] or "").strip()).strip()
            if pc and app:
                by_pc[pc][app].add(release or "Не определена")
                records += 1
    return by_pc, records


def clean(value):
    # Treat data from remote inventory as text, not an Excel formula.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def sheet(book, title, headings, rows, widths):
    tab = book.create_sheet(title)
    tab.append(headings)
    for row in rows:
        tab.append([clean(value) for value in row])
    tab.freeze_panes = "B2"
    tab.auto_filter.ref = tab.dimensions
    tab.sheet_view.showGridLines = False
    tab.row_dimensions[1].height = 30
    for cell in tab[1]:
        cell.fill = PatternFill("solid", fgColor="203864")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for column, width in enumerate(widths, 1):
        tab.column_dimensions[get_column_letter(column)].width = width
    return tab


def main():
    parser = argparse.ArgumentParser(description="Read-only KSC CPU inventory + CAD Excel")
    parser.add_argument("--url", default="https://127.0.0.1:13299")
    parser.add_argument("--insecure-local-test", action="store_true")
    parser.add_argument("--ca-bundle", type=Path)
    parser.add_argument("--cad-csv", type=Path, default=Path("cad_hosts.csv"))
    parser.add_argument("--output", type=Path, default=Path("ksc_hardware_cad_physical.xlsx"))
    parser.add_argument("--probe", action="store_true", help="Show device/VM type coverage before exporting")
    parser.add_argument("--physical-overrides", type=Path, default=None,
                        help="Local CSV (Computer,VerifiedPhysical) confirmed by independent endpoint checks")
    args = parser.parse_args()

    if args.insecure_local_test and args.ca_bundle:
        parser.error("Choose --insecure-local-test OR --ca-bundle")
    if args.insecure_local_test:
        if not args.url.startswith(("https://127.0.0.1:", "https://localhost:")):
            parser.error("The insecure mode is allowed only on a loopback SSH tunnel")
        from urllib3.exceptions import InsecureRequestWarning
        warnings.filterwarnings("once", category=InsecureRequestWarning)
        verify = False
    elif args.ca_bundle:
        if not args.ca_bundle.is_file():
            parser.error("CA bundle does not exist")
        verify = str(args.ca_bundle)
    else:
        parser.error("Specify --insecure-local-test for pilot or --ca-bundle for production")

    cad, cad_count = read_cad(args.cad_csv)
    user = input("KSC username [ksc-inventory]: ").strip() or "ksc-inventory"
    password = getpass.getpass("KSC password: ")
    server = KlAkAdmServer.Create(args.url, user, password, domain="", internal_user=True, verify=verify)
    del password

    hosts = fetch_hosts(server)
    hardware, total, types = fetch_hardware(server)
    if not hosts or not hardware:
        raise RuntimeError("No hosts or hardware records returned")
    print("KSC hosts:", len(hosts), "| Hardware records:", total)
    print("dev_type count:", dict(sorted(types.items(), key=lambda x: str(x[0]))))
    print("HST_VM_TYPE count:", dict(Counter(str(h["vm_type"]) for h in hosts.values())))
    if args.probe:
        for kind in (1, 2, 3, 4):
            print(f"dev_type={kind} records={types[kind]}; hosts={sum(any(d['type']==kind for d in rows) for rows in hardware.values())}")
        return

    verified_physical = set()
    if args.physical_overrides:
        if not args.physical_overrides.is_file():
            parser.error("Missing override CSV: " + str(args.physical_overrides))
        with args.physical_overrides.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            if not {"Computer", "VerifiedPhysical"}.issubset(reader.fieldnames or []):
                parser.error("Override CSV columns must be: Computer,VerifiedPhysical")
            known_names = {item["name"] for item in hosts.values()}
            for item in reader:
                name = (item["Computer"] or "").strip().upper()
                ok = (item["VerifiedPhysical"] or "").strip().lower()
                if ok not in ("yes", "true", "1", ""):
                    parser.error(f"Unexpected VerifiedPhysical flag for {name}: {ok}")
                if not name or ok == "":
                    continue
                if name not in known_names:
                    parser.error(f"Override PC does not exist in current KSC hosts: {name}")
                verified_physical.add(name)
        print("Independently verified physical overrides:", len(verified_physical))

    confirmed, excluded = {}, []
    vm_confirmed = 0
    review_required = 0
    for host_id, h in hosts.items():
        status, reason = host_status(h, hardware.get(host_id, []), verified_physical=h["name"] in verified_physical)
        if status == "Физический ПК":
            confirmed[host_id] = h
        else:
            excluded.append([h["name"], str(h["vm_type"]), status, reason])
            if status == "Виртуальная машина":
                vm_confirmed += 1
            else:
                review_required += 1

    hwrows, details_ram, details_gpu, details_disk = [], [], [], []
    by_pc, model_hosts, cpu_duplicate_hosts = {}, defaultdict(set), 0
    disk_by_pc = {}
    cpu_audit_by_pc = {}
    for host_id, host in sorted(confirmed.items(), key=lambda x: x[1]["name"]):
        pc = host["name"]
        data, cpus, rams, gpus, disks, duplicates = hardware_info(hardware.get(host_id, []))
        cpu_duplicate_hosts += bool(duplicates)
        full = [pc, *data[:-1], host["last_seen"], data[-1]]
        raw_cores, _norm_cores, audit = normalized_cpu_cores(cpus)
        hwrows.append([*full, raw_cores, audit])
        cpu_audit_by_pc[pc] = (raw_cores, audit)
        by_pc[pc] = full
        disk_by_pc[pc] = [type_of_storage(d) for d in disks]
        for cpu in cpus:
            model_hosts[(vendor(cpu["model"]), cpu["model"])].add(pc)
        for r in rams:
            details_ram.append([pc, r["model"], r["manufacturer"], gib(r["xpar3"]), r["ram_type"], r["ram_part"], r["xpar2"]])
        for r in gpus:
            details_gpu.append([pc, r["model"], r["manufacturer"], vram_mb(r["vram"])])
        for r in disks:
            cls, basis = classify_storage(r["model"], r["rpm"])
            details_disk.append([pc, r["model"], r["manufacturer"], gib(r["xpar2"]), cls, basis, r["rpm"]])

    cad_rows, cad_unmatched = [], []
    for pc, installed in sorted(cad.items()):
        if pc not in by_pc:
            cad_unmatched.append(pc)
            if not any(row[0] == pc for row in excluded):
                excluded.append([pc, "", "Проверить вручную", "Инженерный ПК вне подтверждённой физической выборки"])
            continue
        row = by_pc[pc]
        cad_rows.append([*row[:-2],
                         ", ".join(sorted(installed.get("SOLIDWORKS", set()))),
                         ", ".join(sorted(installed.get("КОМПАС-3D", set()))),
                         ", ".join(sorted(installed.get("Siemens NX", set()))),
                         row[-1], *cpu_audit_by_pc[pc]])

    priority_rows = []
    for pc, installed in sorted(cad.items()):
        if pc not in by_pc:
            priority_rows.append([pc, "Проверить", 0, "", "", "", "",
                                  "; ".join(f"{app}: {', '.join(sorted(v))}" for app, v in installed.items()),
                                  "Нет подтверждённой физической записи KSC"])
            continue
        row = by_pc[pc]
        sw = ", ".join(sorted(installed.get("SOLIDWORKS", set())))
        compas = ", ".join(sorted(installed.get("КОМПАС-3D", set())))
        nx = ", ".join(sorted(installed.get("Siemens NX", set())))
        cad_text = "; ".join(p for p in (
            "SOLIDWORKS " + sw if sw else "",
            "КОМПАС-3D " + compas if compas else "",
            "Siemens NX " + nx if nx else "",
        ) if p)
        level, points, reason = modernization_priority(
            row[1], row[8], row[10], disk_by_pc.get(pc, []), cad_text
        )
        priority_rows.append([pc, level, points, row[1], row[8], row[10],
                              row[14], cad_text, reason])
    priority_order = {"Высокий": 0, "Средний": 1, "Низкий": 2, "Проверить": 3}
    priority_rows.sort(key=lambda r: (priority_order.get(r[1], 9), -r[2], r[0]))

    headers = ["Компьютер", "CPU модель", "Производитель", "CPU сырых",
               "CPU уникальных", "Ядра", "Потоки", "Частота CPU (МГц)",
               "RAM (ГиБ)", "Модулей RAM", "GPU модель", "VRAM (МБ)",
               "Диски (ГиБ)", "Накопителей", "Тип дисков",
               "Модели накопителей", "Связь агента KSC", "Контроль"]
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Итоги", ["Показатель", "Значение"], [
        ["Дата", datetime.now().astimezone().isoformat(timespec="seconds")],
        ["KSC устройств", len(hosts)], ["Подтверждённые физические", len(confirmed)],
        ["Исключены: подтверждённые ВМ", vm_confirmed],
        ["Исключены: требуют ручной проверки", review_required],
        ["Устройств с дублями CPU", cpu_duplicate_hosts],
        ["Накопителей с неопределённым классом", sum(r[4] == "Не определён" for r in details_disk)],
        ["GPU с VRAM без сведений", sum(r[3] is None for r in details_gpu)],
        ["CAD компьютеров", len(cad)], ["CAD компьютеров в физической выборке", len(cad_rows)],
        ["CAD компьютеров вне выборки", len(cad_unmatched)],
        ["Критерий", "HST_VM_TYPE=2, нет признаков виртуализации"],
        ["Неизвестный тип", "Исключается до ручной проверки"],
        ["Ручное подтверждение физических", len(verified_physical)],
        ["Пояснение", "Приоритет модернизации — предварительная эвристика, не benchmark"],
        ["Источник", "Сохранённая инвентаризация KSC, не live-проверка"]
    ], [45, 80])
    sheet(wb, "Оборудование", [*headers, "Ядра KSC исходно", "Нормализация ядер"], hwrows,
          [26, 59, 20, 16, 18, 14, 15, 23, 16, 16, 59, 18, 18, 17, 29, 58, 34, 35, 21, 74])
    summary_rows = [[v, model, len(pcs)] for (v, model), pcs in
                    sorted(model_hosts.items(), key=lambda x: (-len(x[1]), x[0][1]))]
    sheet(wb, "Сводка CPU", ["Производитель", "Модель", "Компьютеров"], summary_rows, [20, 65, 19])
    sheet(wb, "CPU + CAD", [*headers[:-2], "SOLIDWORKS", "КОМПАС-3D", "Siemens NX", "Контроль", "Ядра KSC исходно", "Нормализация ядер"],
          cad_rows, [26, 59, 20, 16, 18, 14, 15, 23, 16, 16, 59, 18, 18, 17, 29, 58, 30, 30, 30, 35, 21, 74])
    sheet(wb, "Модули RAM", ["Компьютер", "RAM", "Производитель", "ГиБ", "Тип", "PartNum", "МГц"],
          details_ram, [26, 56, 25, 16, 20, 33, 18])
    sheet(wb, "GPU", ["Компьютер", "Модель GPU", "Производитель", "VRAM МБ (пусто=нет данных)"],
          details_gpu, [26, 59, 26, 26])
    sheet(wb, "Накопители", ["Компьютер", "Модель", "Производитель", "ГиБ", "Класс", "Основание", "RPM"],
          details_disk, [26, 59, 26, 18, 31, 41, 16])
    sheet(wb, "Исключения", ["Компьютер", "HST_VM_TYPE", "Категория", "Причина"],
          sorted(excluded), [28, 21, 29, 69])
    sheet(wb, "Приоритет CAD",
          ["Компьютер", "Приоритет", "Баллы", "CPU", "RAM ГиБ", "GPU", "Диски",
           "Установленный CAD", "Основания"],
          priority_rows, [27, 19, 13, 60, 16, 60, 36, 63, 89])
    sheet(wb, "Методика", ["Правило", "Интерпретация"], [
        ["Приоритет", "Только предварительный список для проверки, не измерение производительности"],
        ["Высокий / Средний / Низкий", "Баллы: >=5 / 2–4 / 0–1; при неизвестном CPU/RAM = Проверить"],
        ["Старый CPU", "Intel Core поколение <=4, Pentium, Celeron, AMD FX/Athlon II/E2: +4"],
        ["Другой старый CPU", "Intel Core поколение 5–7 или Ryzen первого поколения: +2"],
        ["RAM", "<=8 ГиБ: +2; 9–12 ГиБ: +1"],
        ["GPU", "Старые GeForce GT/GTS/GTX и Quadro из явного перечня: +2"],
        ["Накопители", "Только подтверждённые HDD: +1; неизвестный тип не влияет на баллы"],
        ["CAD и CPU", "SOLIDWORKS 2023 или Siemens NX на старой платформе: +1"],
        ["VRAM", "KSC сообщил 0 либо не передал значение -> пусто (не 0 МБ)"],
        ["SSD/HDD", "Модельное семейство/явный тип/RPM; неизвестным оставляем неизвестный"],
        ["VM", "Подтверждённые ВМ и неоднозначные HST_VM_TYPE разделяются"],
        ["Подтверждённый физический ПК", "Только локальный список --physical-overrides после независимой проверки; сигнатуры ВМ имеют приоритет"],
        ["CPU дубликаты", "Одинаковая сигнатура исключена из суммы; реальные два сокета проверить"],
        ["AMD FX-6100, FX-8320", "При сигнатуре KSC (3/6 и 4/8): нормализовать ядра до 6 и 8; исходное поле хранится отдельно"],
        ["Ограничение", "Не учитываются тесты производительности CAD, размеры сборок, цена обновления"],
        ["Источник", "Последняя сохранённая аппаратная инвентаризация KSC"],
    ], [38, 112])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print("===== RESULTS =====")
    print("KSC hosts:", len(hosts), "confirmed physical:", len(confirmed), "excluded/unknown:", len(hosts) - len(confirmed))
    print("Duplicate CPU hosts:", cpu_duplicate_hosts)
    print("Normalized CPU core hosts:", sum(bool(note and "нормализовано" in note) for _,note in cpu_audit_by_pc.values()))
    print("VM signals:", vm_confirmed, "manual-review devices:", review_required)
    print("Physical overrides applied:", len(verified_physical))
    print("Priority CAD:", dict(Counter(row[1] for row in priority_rows)))
    print("RAM rows:", len(details_ram), "GPU rows:", len(details_gpu), "storage rows:", len(details_disk))
    print("CAD computers:", len(cad), "physical with inventory:", len(cad_rows), "excluded/unknown:", len(cad_unmatched))
    if cad_unmatched:
        print("REVIEW excluded CAD machines:", ", ".join(cad_unmatched))
    print("Saved:", args.output.resolve())



if __name__ == "__main__":
    main()
