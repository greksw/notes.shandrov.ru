#!/usr/bin/env python3
"""Read-only inventory of installed CAD software via KSC 15.1 Open API.

Uses the KlAkOAPI library distributed with Kaspersky Security Center.
This is a sanitized example adapted from a verified KSC 15.1 inventory run.
Validate the public version against your environment before scheduling.
"""
import argparse
import csv
import getpass
import re
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from KlAkOAPI.AdmServer import KlAkAdmServer
from KlAkOAPI.HostGroup import KlAkHostGroup
from KlAkOAPI.ChunkAccessor import KlAkChunkAccessor
from KlAkOAPI.InventoryApi import KlAkInventoryApi
from KlAkOAPI.Params import KlAkParams

HEADERS = ["Computer", "Software", "Release", "ServicePack",
           "DisplayVersion", "DisplayName", "Publisher"]


def prop(item, name):
    if hasattr(item, "GetNames"):
        return item[name] if name in item.GetNames() else None
    if isinstance(item, dict):
        return item.get(name)
    return None


def text_prop(item, name):
    value = prop(item, name)
    return "" if value is None else str(value)


def classify(name):
    """Match primary CAD applications only, not viewers/add-ons/APIs."""
    name = name.strip()
    if re.search(r"viewer|просмотрщик", name, re.I):
        return None
    kompas = re.fullmatch(r"КОМПАС-3D\s+v(\d+(?:\.\d+)?)\s*(?:x64|x86)?", name, re.I)
    if kompas:
        return ("КОМПАС-3D", "v" + kompas.group(1), "")
    sw = re.fullmatch(r"SOLIDWORKS\s+(20\d{2})(?:\s+SP\s*(\d+(?:\.\d+)?))?", name, re.I)
    if sw:
        return ("SOLIDWORKS", sw.group(1), "SP" + sw.group(2) if sw.group(2) else "")
    nx = re.fullmatch(r"(?:Siemens\s+NX|UGS\s+NX|NX)(?:\s+(\d+(?:\.\d+)*))?", name, re.I)
    if nx:
        return ("Siemens NX", nx.group(1) or "", "")
    return None


def products_from(response):
    params = response.RetVal()
    keys = params.GetNames()
    for key in ("GNRL_EA_PARAM_1", "KLEVP_EA_PARAM_1"):
        if key in keys:
            return params[key]
    raise RuntimeError("Unexpected inventory response fields: " + repr(list(keys)))


def list_hosts(server):
    api = KlAkHostGroup(server)
    chunks = KlAkChunkAccessor(server)
    response = api.FindHosts(
        '(KLHST_WKS_HOSTNAME="*")',
        ["KLHST_WKS_HOSTNAME", "KLHST_WKS_DN"],
        [],
        {"KLGRP_FIND_FROM_CUR_VS_ONLY": True},
        3600,
    )
    accessor = response.OutPar("strAccessor")
    hosts = []
    try:
        total = chunks.GetItemsCount(accessor).RetVal()
        for offset in range(0, total, 100):
            result = chunks.GetItemsChunk(accessor, offset, min(100, total - offset))
            entries = result.OutPar("pChunk")["KLCSP_ITERATOR_ARRAY"]
            for item in entries:
                # In the verified KSC 15.1 response this field was used as host ID.
                host_id = text_prop(item, "KLHST_WKS_HOSTNAME")
                name = text_prop(item, "KLHST_WKS_DN")
                if host_id:
                    hosts.append((host_id, name or host_id))
    finally:
        chunks.Release(accessor)
    return hosts


def csv_report(path, rows, headers):
    with path.open("w", encoding="utf-8-sig", newline="") as out:
        writer = csv.DictWriter(out, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def sheet(workbook, title, headers, rows):
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="203864")
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        width = max((len(str(cell.value or "")) for cell in column), default=12) + 2
        ws.column_dimensions[letter].width = min(max(width, 14), 64)
    return ws


def xlsx_report(path, rows, checked, errors):
    wb = Workbook()
    wb.remove(wb.active)
    by_host = defaultdict(lambda: defaultdict(set))
    by_version = defaultdict(set)
    for row in rows:
        app = row["Software"]
        pc = row["Computer"]
        label = " ".join(x for x in (row["Release"], row["ServicePack"]) if x)
        if not label:
            label = "build " + row["DisplayVersion"] if row["DisplayVersion"] else "unknown release"
        by_host[pc][app].add(label)
        key = (app, row["Release"], row["ServicePack"], row["DisplayVersion"])
        by_version[key].add(pc)

    apps = ("SOLIDWORKS", "КОМПАС-3D", "Siemens NX")
    sheet(wb, "По компьютерам", ["Computer", *apps], [
        [pc, *[", ".join(sorted(apps_at_pc.get(app, set()))) for app in apps]]
        for pc, apps_at_pc in sorted(by_host.items())
    ])
    sheet(wb, "Все установки", HEADERS, [[row[h] for h in HEADERS] for row in rows])
    for app in apps:
        sheet(wb, app, HEADERS, [[r[h] for h in HEADERS] for r in rows if r["Software"] == app])
    sheet(wb, "Сводка по версиям",
          ["ПО", "Выпуск", "SP", "Build", "Компьютеров"],
          [[*key, len(pcs)] for key, pcs in sorted(by_version.items())])
    sheet(wb, "Контроль", ["Показатель", "Значение"], [
        ["Сформировано", datetime.now().astimezone().isoformat(timespec="seconds")],
        ["Компьютеров проверено", checked],
        ["CAD установок", len(rows)],
        ["Уникальных компьютеров с CAD", len(by_host)],
        ["Ошибок запросов API", len(errors)],
        ["Источник", "Последняя сохранённая инвентаризация KSC, не live-проверка"],
        ["Связь с агентами", "В этой версии скрипта не проверяется"],
    ])
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description="KSC CAD inventory (read-only)")
    parser.add_argument("--url", default="https://127.0.0.1:13299")
    parser.add_argument("--ca-bundle", type=Path, help="Trust chain for a certificate-valid API hostname")
    parser.add_argument("--insecure-local-test", action="store_true",
                        help="Temporary test only: SSH forwarding on 127.0.0.1, without certificate checks")
    parser.add_argument("--limit", type=int, default=0, help="0: all KSC hosts; positive: first N")
    parser.add_argument("--output-dir", type=Path, default=Path("reports"))
    args = parser.parse_args()

    if args.limit < 0:
        parser.error("--limit must be >= 0")
    if args.insecure_local_test and args.ca_bundle:
        parser.error("Select only one TLS verification mode")
    if args.insecure_local_test:
        if not (args.url.startswith("https://127.0.0.1:") or
                args.url.startswith("https://localhost:")):
            parser.error("--insecure-local-test requires a loopback SSH tunnel")
        from urllib3.exceptions import InsecureRequestWarning
        warnings.filterwarnings("once", category=InsecureRequestWarning)
        verify = False
    elif args.ca_bundle:
        if not args.ca_bundle.is_file():
            parser.error("CA bundle file does not exist")
        verify = str(args.ca_bundle)
    else:
        parser.error("Specify --ca-bundle (production) or --insecure-local-test (pilot)")

    username = input("KSC internal user: ").strip()
    if not username:
        parser.error("KSC username must not be empty")
    password = getpass.getpass("KSC password: ")
    server = KlAkAdmServer.Create(
        args.url, username, password, domain="", internal_user=True, verify=verify
    )
    del password
    inventory = KlAkInventoryApi(server)
    hosts = list_hosts(server)
    print("Computers found:", len(hosts))
    if not hosts:
        raise RuntimeError("No computers returned by KSC; refusing to emit a misleading empty report")
    if args.limit:
        hosts = hosts[:args.limit]

    rows, errors, seen = [], [], set()
    for n, (host_id, computer) in enumerate(hosts, 1):
        try:
            for p in products_from(inventory.GetHostInvProducts(host_id, KlAkParams())):
                name = text_prop(p, "DisplayName")
                classification = classify(name)
                if classification is None:
                    continue
                software, release, sp = classification
                build = text_prop(p, "DisplayVersion")
                identity = (host_id, name, build)
                if identity in seen:
                    continue
                seen.add(identity)
                rows.append(dict(
                    Computer=computer, Software=software, Release=release,
                    ServicePack=sp, DisplayVersion=build, DisplayName=name,
                    Publisher=text_prop(p, "Publisher")
                ))
        except Exception as exc:
            errors.append(dict(Computer=computer, Error=f"{type(exc).__name__}: {exc}"))
        if n % 25 == 0 or n == len(hosts):
            print(f"Processed {n}/{len(hosts)}; CAD installations: {len(rows)}; API errors: {len(errors)}")

    rows.sort(key=lambda r: (r["Computer"].casefold(), r["Software"], r["Release"], r["ServicePack"]))
    args.output_dir.mkdir(parents=True, exist_ok=True)
    csv_report(args.output_dir / "cad_hosts.csv", rows, HEADERS)
    csv_report(args.output_dir / "cad_errors.csv", errors, ["Computer", "Error"])
    xlsx_report(args.output_dir / "cad_hosts.xlsx", rows, len(hosts), errors)
    print(f"Done: {len(rows)} installations on {len({r['Computer'] for r in rows})} computers; errors: {len(errors)}")
    print("Output:", args.output_dir.resolve())
    if errors:
        print("WARNING: some host inventory requests failed; check cad_errors.csv.")


if __name__ == "__main__":
    main()
